import os
from typing import Tuple, List
import PyPDF2
from docx import Document as DocxDocument


def extract_text_from_pdf(file_path: str) -> Tuple[str, str]:
    """
    Extract text from a PDF file.
    Returns (filename, extracted_text).
    """
    text_content = []
    
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                text_content.append(page.extract_text())
    except Exception as e:
        raise ValueError(f"Error reading PDF: {str(e)}")
    
    filename = os.path.basename(file_path)
    extracted_text = "\n".join(text_content)
    
    return filename, extracted_text


def extract_text_from_docx(file_path: str) -> Tuple[str, str]:
    """
    Extract text from a DOCX file.
    Returns (filename, extracted_text).
    """
    try:
        doc = DocxDocument(file_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_content.append(cell.text)
    except Exception as e:
        raise ValueError(f"Error reading DOCX: {str(e)}")
    
    filename = os.path.basename(file_path)
    extracted_text = "\n".join(text_content)
    
    return filename, extracted_text


def extract_text_from_txt(file_path: str) -> Tuple[str, str]:
    """
    Extract text from a TXT file.
    Returns (filename, extracted_text).
    """
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
            text_content = file.read()
    except Exception as e:
        raise ValueError(f"Error reading TXT: {str(e)}")
    
    filename = os.path.basename(file_path)
    return filename, text_content


def extract_text_from_document(file_path: str) -> Tuple[str, str]:
    """
    Extract text from a document based on file extension.
    Supports PDF, DOCX, and TXT files.
    Returns (filename, extracted_text).
    """
    file_extension = os.path.splitext(file_path)[1].lower()
    
    if file_extension == '.pdf':
        return extract_text_from_pdf(file_path)
    elif file_extension == '.docx':
        return extract_text_from_docx(file_path)
    elif file_extension == '.txt':
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_extension}. Supported formats: PDF, DOCX, TXT")


def chunk_document_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[str]:
    """
    Split document text into overlapping chunks suitable for embeddings.
    """
    chunks = []
    for i in range(0, len(text), chunk_size - chunk_overlap):
        chunk = text[i : i + chunk_size]
        if chunk.strip():
            chunks.append(chunk)
    return chunks
